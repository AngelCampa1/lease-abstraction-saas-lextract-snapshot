"""Tests for canonical Lextract brand asset usage."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock, patch

from docx import Document
from openpyxl import load_workbook

from app.core.config import get_settings
from app.services.exports.excel import ExcelExporter
from app.services.exports.word import WordExporter

SAMPLE_EXTRACTION = {
    "landlord_legal_name": {"value": "ABC Corp"},
    "tenant_legal_name": {"value": "Acme Inc"},
    "premises_address": {"value": "123 Main St"},
    "base_rent_annual": {"value": 120000},
    "lease_term_months": {"value": 60},
}

SAMPLE_CONFIDENCE = {
    "landlord_legal_name": {"score": 0.95, "tier": "high"},
    "tenant_legal_name": {"score": 0.88, "tier": "high"},
}


def test_email_templates_include_canonical_logo_url() -> None:
    from app.services.email import EmailService
    from app.services.brand import BRAND_ASSETS

    svc = EmailService(api_key="re_test_key")
    with patch("app.services.email.resend") as mock_resend:
        mock_resend.Emails.send = MagicMock(return_value={"id": "msg-123"})

        svc.send_extraction_complete(
            to="user@example.com",
            document_name="lease.pdf",
            field_count=126,
            confidence_summary="95%",
            results_url="https://lextract.io/results/ext-1",
        )

    payload = mock_resend.Emails.send.call_args[0][0]
    assert BRAND_ASSETS.email_logo_url in payload["html"]
    assert 'alt="Lextract"' in payload["html"]


def test_brand_asset_urls_use_configured_frontend_url(monkeypatch) -> None:
    from app.services.brand import BRAND_ASSETS

    get_settings.cache_clear()
    monkeypatch.setenv("FRONTEND_URL", "https://preview.lextract.test/")

    try:
        assert (
            BRAND_ASSETS.email_logo_url
            == "https://preview.lextract.test/brand/lextract-email-logo.png"
        )
        assert (
            BRAND_ASSETS.logo_url
            == "https://preview.lextract.test/brand/lextract-logo.png"
        )
        assert (
            BRAND_ASSETS.icon_url
            == "https://preview.lextract.test/brand/lextract-icon.png"
        )
    finally:
        get_settings.cache_clear()


def test_email_templates_include_visible_logo_text_fallback() -> None:
    template_root = Path("app/services/email_templates")
    template_paths = [
        template_root / "cam_flags_found.html",
        template_root / "extraction_complete.html",
        *sorted((template_root / "nurture").glob("*.html")),
    ]

    for template_path in template_paths:
        html = template_path.read_text(encoding="utf-8")
        assert 'data-brand-text="Lextract"' in html, template_path


def test_brand_assets_expose_inline_logo_data_uri() -> None:
    from app.services.brand import BRAND_ASSETS

    assert BRAND_ASSETS.logo_path.is_file()
    assert BRAND_ASSETS.logo_data_uri.startswith("data:image/png;base64,")


def test_guest_account_email_includes_canonical_logo_url() -> None:
    from app.services.email import EmailService
    from app.services.brand import BRAND_ASSETS

    svc = EmailService(api_key="re_test_key")
    with patch("app.services.email.resend") as mock_resend:
        mock_resend.Emails.send = MagicMock(return_value={"id": "msg-456"})

        svc.send_complete_your_account(
            to="guest@example.com",
            results_url="https://lextract.io/results/ext-1",
            password_reset_url="https://lextract.io/reset",
        )

    payload = mock_resend.Emails.send.call_args[0][0]
    assert BRAND_ASSETS.email_logo_url in payload["html"]
    assert 'alt="Lextract"' in payload["html"]
    assert 'data-brand-text="Lextract"' in payload["html"]


def test_pdf_export_cover_includes_logo_image() -> None:
    from app.services.brand import BRAND_ASSETS

    try:
        from app.services.exports.pdf import PdfExporter
    except OSError:
        import pytest

        pytest.skip("WeasyPrint system dependencies are unavailable")

    html = PdfExporter()._build_html(
        extraction_data=SAMPLE_EXTRACTION,
        confidence_scores=SAMPLE_CONFIDENCE,
        red_flags=[],
        template="commercial",
        document_filename="lease.pdf",
    )

    assert BRAND_ASSETS.logo_data_uri in html
    assert 'alt="Lextract"' in html


def test_word_export_cover_embeds_logo_image() -> None:
    doc_bytes = WordExporter().generate(
        extraction_data=SAMPLE_EXTRACTION,
        confidence_scores=SAMPLE_CONFIDENCE,
        red_flags=[],
        template="commercial",
        document_filename="lease.pdf",
    )

    doc = Document(BytesIO(doc_bytes))
    assert len(doc.inline_shapes) >= 1


def test_excel_summary_embeds_logo_image() -> None:
    workbook_bytes = ExcelExporter().generate(
        extraction_data=SAMPLE_EXTRACTION,
        confidence_scores=SAMPLE_CONFIDENCE,
        red_flags=[],
        template="commercial",
        document_filename="lease.pdf",
    )

    workbook = load_workbook(BytesIO(workbook_bytes))
    summary = workbook["Summary"]
    assert len(summary._images) >= 1
